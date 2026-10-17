# -*- coding: utf-8 -*-
"""锚定式就地改稿工具库（保留原排版）。配套 5 份文档的整改脚本使用。

约定：
  * 改稿脚本结构：load() 在顶部 → 全部编辑 → save() 在最后一行。
    锚点写错会在写入前抛错，整次运行等于空操作，可安全重跑。
  * doc.paragraphs 看不到表格单元格内的段落 → 表格一律用 tbl_after / row_of 定位。
  * 锚点必须原文照抄（含全角标点）。
"""
import copy

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph
from lxml import etree


# ---------------------------------------------------------------- load / save
def load(path):
    return Document(path)


def save(doc, path):
    doc.save(path)


# ---------------------------------------------------------------- paragraphs
def find_p(doc, needle, nth=0, must=True):
    hits = [p for p in doc.paragraphs if needle in p.text]
    if len(hits) <= nth:
        if must:
            raise LookupError('paragraph not found: %r (hits=%d)' % (needle, len(hits)))
        return None
    return hits[nth]


def set_text(p, text):
    runs = p.runs
    if runs:
        keep = runs[0]
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
        keep.text = text
    else:
        p.add_run(text)
    return p


def replace_in(p, old, new):
    if old not in p.text:
        raise LookupError('%r not in %r' % (old, p.text[:90]))
    return set_text(p, p.text.replace(old, new))


def edit(doc, needle, old, new, nth=0):
    return replace_in(find_p(doc, needle, nth), old, new)


def insert_after(p, lines, template=None):
    tpl = template or p
    anchor = p._p
    out = []
    for line in lines:
        new = copy.deepcopy(tpl._p)
        anchor.addnext(new)
        anchor = new
        np = Paragraph(new, tpl._parent)
        set_text(np, line)
        out.append(np)
    return out


def insert_before(p, lines, template=None):
    tpl = template or p
    out = []
    for line in lines:
        new = copy.deepcopy(tpl._p)
        p._p.addprevious(new)
        np = Paragraph(new, tpl._parent)
        set_text(np, line)
        out.append(np)
    return out


def delete_p(p):
    p._p.getparent().remove(p._p)


def replace_section(doc, start_needle, end_needle, lines):
    """替换两个锚点之间的全部段落（保留两端锚点本身），以 start 的下一段为格式模板。"""
    ps = doc.paragraphs
    i0 = next(i for i, p in enumerate(ps) if start_needle in p.text)
    i1 = next(i for i, p in enumerate(ps) if end_needle in p.text and i > i0)
    tpl = ps[i0 + 1]
    for p in ps[i0 + 1:i1]:
        delete_p(p)
    return insert_after(ps[i0], lines, template=tpl)


# ---------------------------------------------------------------- tables
def body_children(doc):
    out = []
    for el in doc.element.body:
        tag = etree.QName(el).localname
        if tag == 'p':
            out.append(('p', Paragraph(el, doc)))
        elif tag == 'tbl':
            out.append(('tbl', el))
    return out


def tbl_after(doc, heading_text, nth=0):
    kids = body_children(doc)
    seen = 0
    for i, (kind, obj) in enumerate(kids):
        if kind == 'p' and heading_text in obj.text:
            if seen < nth:
                seen += 1
                continue
            for kind2, obj2 in kids[i + 1:]:
                if kind2 == 'tbl':
                    return Table(obj2, doc)
                if kind2 == 'p' and (obj2.style is not None and obj2.style.name.startswith('Heading')):
                    break
    raise LookupError('no table after heading %r' % heading_text)


def row_of(table, key):
    """按首列文本前缀定位行号（表格增删行后仍稳定）。"""
    for i, row in enumerate(table.rows):
        if row.cells[0].text.strip().startswith(key):
            return i
    raise LookupError('row not found in table: %r' % key)


def set_cell(cell, text):
    ps = cell.paragraphs
    set_text(ps[0], text)
    for p in ps[1:]:
        delete_p(p)
    return cell


def add_row_after(table, idx, values, template_idx=None):
    tpl = table.rows[template_idx if template_idx is not None else idx]._tr
    new = copy.deepcopy(tpl)
    tpl.addnext(new)
    row = _Row(new, table)
    for c, v in zip(row.cells, values):
        set_cell(c, v)
    return row


def delete_row(table, idx):
    tr = table.rows[idx]._tr
    tr.getparent().remove(tr)


# ---------------------------------------------------------------- new table
def make_table(doc, anchor_p, rows_data, widths_frac, tpl_table, after=True):
    """按模板表的内联格式新建一张表，插在 anchor_p 之后/之前。

    rows_data: [[c1, c2, ...], ...]；widths_frac: 各列宽度占比（合计 1.0）。
    """
    tbl = OxmlElement('w:tbl')
    tpr = tpl_table._tbl.find(qn('w:tblPr'))
    if tpr is not None:
        tbl.append(copy.deepcopy(tpr))
    sec = doc.sections[0]
    total = int((sec.page_width - sec.left_margin - sec.right_margin) / 635)
    grid = OxmlElement('w:tblGrid')
    for f in widths_frac:
        gc = OxmlElement('w:gridCol')
        gc.set(qn('w:w'), str(int(total * f)))
        grid.append(gc)
    tbl.append(grid)
    tpl_par = copy.deepcopy(tpl_table.rows[1].cells[0].paragraphs[0]._p)
    for rowdata in rows_data:
        tr = OxmlElement('w:tr')
        for ci, val in enumerate(rowdata):
            tc = OxmlElement('w:tc')
            tcPr = OxmlElement('w:tcPr')
            tcW = OxmlElement('w:tcW')
            tcW.set(qn('w:w'), str(int(total * widths_frac[ci])))
            tcW.set(qn('w:type'), 'dxa')
            tcPr.append(tcW)
            tc.append(tcPr)
            par = copy.deepcopy(tpl_par)
            tc.append(par)
            tr.append(tc)
        tbl.append(tr)
    if after:
        anchor_p._p.addnext(tbl)
    else:
        anchor_p._p.addprevious(tbl)
    tab = Table(tbl, doc)
    for r, rowdata in enumerate(rows_data):
        for c, val in enumerate(rowdata):
            set_cell(tab.rows[r].cells[c], val)
    return tab


def add_subsection(doc, before_needle, title, intro_lines, table_rows, h_tpl, b_tpl,
                   t_tpl, widths=(0.18, 0.14, 0.14, 0.54)):
    """在 before_needle 段之前插入：标题 + 说明段 + 表格。"""
    stop = find_p(doc, before_needle)
    h = insert_before(stop, [title], template=h_tpl)[0]
    ps = insert_after(h, intro_lines, template=b_tpl)
    return make_table(doc, ps[-1], table_rows, widths, t_tpl)


def add_tail_section(doc, end_needle, title, intro_lines, table_rows, h_tpl, b_tpl, t_tpl,
                     widths=(0.10, 0.18, 0.58, 0.14)):
    """在文末「— 文档结束 —」之前插入：效力顺序说明 + 修订记录表。"""
    return add_subsection(doc, end_needle, title, intro_lines, table_rows, h_tpl, b_tpl,
                          t_tpl, widths=widths)
