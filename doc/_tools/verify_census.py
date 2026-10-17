# -*- coding: utf-8 -*-
"""只读校验：标题普查（逐条列出，便于人眼确认无杂散标题）+ 表格方正 + 段落/表格计数。"""
import glob
import os
import sys

from docx import Document


def main(src):
    for f in sorted(glob.glob(os.path.join(src, '*.docx'))):
        if os.path.basename(f).startswith('~$'):
            continue
        d = Document(f)
        heads = [(p.style.name, p.text.strip()) for p in d.paragraphs
                 if p.style is not None and p.style.name.startswith('Heading')]
        h1 = [t for s, t in heads if s in ('Heading 1', 'Heading1')]
        h2 = [t for s, t in heads if s in ('Heading 2', 'Heading2')]
        h3 = [t for s, t in heads if s in ('Heading 3', 'Heading3')]
        bad = []
        for i, t in enumerate(d.tables, 1):
            widths = {len(r.cells) for r in t.rows}
            if len(widths) > 1:
                bad.append((i, sorted(widths)))
        print('=' * 8, os.path.basename(f)[:34])
        print('  段落=%d 表格=%d | H1=%d H2=%d H3=%d | 非方正表=%s'
              % (len(d.paragraphs), len(d.tables), len(h1), len(h2), len(h3), bad or '无'))
        print('  H1:', ' / '.join(h1))
        print('  H2:', ' / '.join(h2))
        if h3:
            print('  H3:', ' / '.join(h3))
        suspects = [t for t in h1 + h2 if not (t[:1].isdigit() or '媒体下载器' in t)]
        if suspects:
            print('  !! 疑似杂散标题:', suspects)
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1]))
